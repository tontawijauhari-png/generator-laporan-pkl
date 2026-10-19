import io
import json
import re
from datetime import datetime
from pydantic import BaseModel, Field
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from google import genai
from google.genai import types

st.set_page_config(
    page_title="Generator Laporan & Logbook PKL SMK",
    page_icon="📝",
    layout="wide"
)

# --- SKEMA OUTPUT JSON DARI AI ---
class LogbookEntry(BaseModel):
    minggu_ke: int = Field(description="Nomor urut minggu ke- (1 sampai N)")
    uraian_pekerjaan: str = Field(description="Uraian detail tugas/kegiatan teknis yang dikerjakan pada minggu tersebut")
    capaian_belajar: str = Field(description="Kompetensi, keterampilan, atau hasil nyata yang dicapai siswa")

class LaporanPKLOutput(BaseModel):
    kerangka_laporan: str = Field(description="Teks kerangka lengkap BAB I sampai BAB IV beserta sub-babnya")
    logbook_mingguan: list[LogbookEntry] = Field(description="Daftar kegiatan mingguan sesuai jumlah minggu yang diminta")


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Mengatur padding/margin internal pada sel tabel."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_toc_line(doc, title, page_str, is_bold=False, indent_left=0.0):
    """
    Menambahkan baris Daftar Isi dengan tab stop titik-titik (dot leader)
    yang rapi sampai margin kanan.
    """
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    if indent_left > 0:
        p.paragraph_format.left_indent = Inches(indent_left)

    # Tambahkan Tab Stop di batas kanan (Inches 5.5 dari margin kiri)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(5.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    run_title = p.add_run(title)
    run_title.bold = is_bold
    run_title.font.size = Pt(11)

    p.add_run("\t")  # Karakter tab yang memicu dot leader

    run_page = p.add_run(page_str)
    run_page.bold = is_bold
    run_page.font.size = Pt(11)


def create_word_document(
    nama_siswa, nis, sekolah, kota, tahun_ajaran, judul_text,
    ai_kerangka, ai_logbook, jurusan, tempat_dudi, durasi_text, total_minggu,
    nama_pembimbing_sekolah, nama_pembimbing_dudi, nama_kepsek,
    logo_file=None
):
    """
    Menyusun dokumen Word lengkap:
    1. Cover Resmi dengan Logo
    2. Lembar Pengesahan Resmi 3 Pihak
    3. Halaman Kata Pengantar
    4. Format Struktur Daftar Isi (TOC) Ber-tabulasi Titik-titik
    5. Kerangka Laporan Bab I - IV
    6. Tabel Logbook Mingguan Terisi Penuh Sesuai Durasi
    """
    doc = Document()

    # --- PENGATURAN HALAMAN (A4 & Margin Standar 4-3-3-3 cm) ---
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.18)    # ~3 cm
    section.bottom_margin = Inches(1.18) # ~3 cm
    section.left_margin = Inches(1.57)   # ~4 cm
    section.right_margin = Inches(1.18)  # ~3 cm

    # --- STYLE DEFAULT: TIMES NEW ROMAN 12pt SPASI 1.5 ---
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(12)
    normal_font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(4)

    # =========================================================================
    # 1. HALAMAN COVER RESMI
    # =========================================================================
    p_cov_hdr = doc.add_paragraph()
    p_cov_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov_hdr.paragraph_format.space_after = Pt(4)
    run_ch1 = p_cov_hdr.add_run("LAPORAN PRAKTIK KERJA LAPANGAN (PKL)\n")
    run_ch1.bold = True
    run_ch1.font.size = Pt(14)

    p_cov_title = doc.add_paragraph()
    p_cov_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov_title.paragraph_format.space_before = Pt(8)
    p_cov_title.paragraph_format.space_after = Pt(12)
    run_ctitle = p_cov_title.add_run(f'"{judul_text.upper()}"\n')
    run_ctitle.bold = True
    run_ctitle.font.size = Pt(13)

    if tempat_dudi:
        run_cdudi = p_cov_title.add_run(f'PADA {tempat_dudi.upper()}')
        run_cdudi.bold = True
        run_cdudi.font.size = Pt(12)

    p_cov_sub = doc.add_paragraph()
    p_cov_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov_sub.paragraph_format.line_spacing = 1.15
    p_cov_sub.paragraph_format.space_before = Pt(6)
    p_cov_sub.paragraph_format.space_after = Pt(16)
    run_csub = p_cov_sub.add_run(
        f"Diajukan sebagai salah satu syarat untuk menyelesaikan\n"
        f"Praktik Kerja Lapangan ({durasi_text}) dan Uji Kompetensi Keahlian"
    )
    run_csub.font.size = Pt(10)
    run_csub.italic = True

    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(10)
    p_logo.paragraph_format.space_after = Pt(16)

    if logo_file is not None:
        logo_file.seek(0)
        p_logo.add_run().add_picture(logo_file, width=Inches(2.2))
    else:
        r_box = p_logo.add_run(
            "┌────────────────────────────────┐\n"
            "│                                │\n"
            "│       [ TEMPEL LOGO RESMI ]    │\n"
            "│          SEKOLAH DI SINI       │\n"
            "│                                │\n"
            "└────────────────────────────────┘"
        )
        r_box.font.size = Pt(10)
        r_box.font.color.rgb = RGBColor(128, 128, 128)

    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author.paragraph_format.line_spacing = 1.2
    p_author.paragraph_format.space_before = Pt(10)
    p_author.paragraph_format.space_after = Pt(20)

    r_oleh = p_author.add_run("Disusun Oleh:\n")
    r_oleh.font.size = Pt(11)
    
    r_nama = p_author.add_run(f"{nama_siswa.upper() if nama_siswa else 'NAMA SISWA'}\n")
    r_nama.bold = True
    r_nama.font.size = Pt(12)

    r_nis = p_author.add_run(f"NIS/NISN: {nis if nis else '..............................'}")
    r_nis.font.size = Pt(11)

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.line_spacing = 1.15
    p_inst.paragraph_format.space_before = Pt(16)

    r_inst1 = p_inst.add_run(f"KOMPETENSI KEAHLIAN {jurusan.upper()}\n")
    r_inst1.bold = True
    r_inst1.font.size = Pt(12)

    r_inst2 = p_inst.add_run(f"{sekolah.upper() if sekolah else 'SMK NEGERI / SWASTA'}\n")
    r_inst2.bold = True
    r_inst2.font.size = Pt(13)

    lokasi_thn = f"{kota.upper() if kota else 'KOTA/KABUPATEN'}\nTAHUN AJARAN {tahun_ajaran}"
    r_inst3 = p_inst.add_run(lokasi_thn)
    r_inst3.bold = True
    r_inst3.font.size = Pt(11)

    doc.add_page_break()

    # =========================================================================
    # 2. HALAMAN LEMBAR PENGESAHAN
    # =========================================================================
    p_sah_title = doc.add_paragraph()
    p_sah_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sah = p_sah_title.add_run("LEMBAR PENGESAHAN\nLAPORAN PRAKTIK KERJA LAPANGAN (PKL)\n")
    run_sah.bold = True
    run_sah.font.size = Pt(13)

    identitas_data = [
        ("Judul Laporan", f': "{judul_text.upper()}"'),
        ("Nama Siswa", f": {nama_siswa.upper() if nama_siswa else '...................................................'}"),
        ("NIS / NISN", f": {nis if nis else '...................................................'}"),
        ("Kompetensi Keahlian", f": {jurusan}"),
        ("Durasi / Periode PKL", f": {durasi_text} ({total_minggu} Minggu)"),
        ("Tempat PKL / DUDI", f": {tempat_dudi.upper() if tempat_dudi else '...................................................'}"),
        ("Asal Sekolah", f": {sekolah.upper() if sekolah else '...................................................'}")
    ]
    for label, val in identitas_data:
        p_row = doc.add_paragraph()
        p_row.paragraph_format.line_spacing = 1.15
        p_row.paragraph_format.space_after = Pt(2)
        r_lbl = p_row.add_run(f"{label:<22}")
        r_lbl.bold = True
        p_row.add_run(val)

    p_ket = doc.add_paragraph(
        "Laporan ini telah diperiksa, disetujui, dan disahkan pada:\n"
        "Hari / Tanggal : ................................................\n"
        "Tempat         : ................................................"
    )
    p_ket.paragraph_format.line_spacing = 1.15
    p_ket.paragraph_format.space_before = Pt(8)
    p_ket.paragraph_format.space_after = Pt(16)

    tbl_ttd1 = doc.add_table(rows=2, cols=2)
    tbl_ttd1.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ttd1.autofit = False

    col_widths = [Inches(2.7), Inches(2.7)]
    for row in tbl_ttd1.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    c0 = tbl_ttd1.cell(0, 0)
    c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c0.paragraphs[0].add_run("Menyetujui,\nPembimbing DU/DI / Instansi\n\n\n\n\n").bold = True
    c0.paragraphs[0].add_run(f"( {nama_pembimbing_dudi or '...................................................'} )\nNIP/NIK: .......................................")

    c1 = tbl_ttd1.cell(0, 1)
    c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c1.paragraphs[0].add_run("Menyetujui,\nPembimbing Sekolah\n\n\n\n\n").bold = True
    c1.paragraphs[0].add_run(f"( {nama_pembimbing_sekolah or '...................................................'} )\nNIP: ...........................................")

    p_kepsek = doc.add_paragraph()
    p_kepsek.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_kepsek.paragraph_format.space_before = Pt(16)
    p_kepsek.paragraph_format.line_spacing = 1.15
    p_kepsek.add_run("Mengetahui,\nKepala Sekolah\n\n\n\n\n").bold = True
    p_kepsek.add_run(f"( {nama_kepsek or '...................................................'} )\nNIP: ...........................................")

    doc.add_page_break()

    # =========================================================================
    # 3. HALAMAN KATA PENGANTAR
    # =========================================================================
    p_kp_hdr = doc.add_paragraph()
    p_kp_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_kph = p_kp_hdr.add_run("KATA PENGANTAR\n")
    run_kph.bold = True
    run_kph.font.size = Pt(14)

    p_kp1 = doc.add_paragraph()
    p_kp1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kp1.paragraph_format.first_line_indent = Inches(0.4)
    p_kp1.add_run(
        "Puji dan syukur penulis panjatkan kehadirat Tuhan Yang Maha Esa atas segala limpahan rahmat dan karunia-Nya, "
        f"sehingga penulis dapat menyelesaikan kegiatan Praktik Kerja Lapangan (PKL) di {tempat_dudi or 'Dunia Usaha/Dunia Industri'} "
        f"serta menyelesaikan penyusunan laporan yang berjudul \"{judul_text}\" ini dengan baik dan tepat pada waktunya."
    )

    p_kp2 = doc.add_paragraph()
    p_kp2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kp2.paragraph_format.first_line_indent = Inches(0.4)
    p_kp2.add_run(
        "Laporan ini disusun sebagai salah satu syarat kelulusan dalam menempuh Praktik Kerja Lapangan (PKL) "
        f"pada Kompetensi Keahlian {jurusan} di {sekolah or 'SMK'}. Penyusunan laporan ini tidak lepas dari bimbingan, "
        "bantuan, serta motivasi dari berbagai pihak. Oleh karena itu, penulis ingin menyampaikan ucapan terima kasih yang sebesar-besarnya kepada:"
    )

    # Butir Penghargaan / Pihak
    penghargaan_list = [
        f"Bapak/Ibu {nama_kepsek or 'Kepala Sekolah'}, selaku Kepala {sekolah or 'SMK'} yang telah memberikan izin dan dukungan penuh pelaksanaan PKL.",
        f"Pimpinan dan jajaran staf {tempat_dudi or 'perusahaan tempat PKL'} yang telah memberikan kesempatan serta fasilitas belajar selama masa magang.",
        f"Bapak/Ibu {nama_pembimbing_dudi or 'Pembimbing Industri'}, selaku pembimbing di DU/DI yang telah tulus membimbing dan membagikan ilmu teknis di dunia kerja nyata.",
        f"Bapak/Ibu {nama_pembimbing_sekolah or 'Pembimbing Sekolah'}, selaku guru pembimbing yang senantiasa mendampingi dan memberikan arahan hingga laporan ini selesai.",
        f"Bapak dan Ibu Dewan Guru serta Staf Tata Usaha {sekolah or 'SMK'} atas bekal ilmu dan fasilitas yang telah diberikan.",
        "Kedua orang tua dan segenap keluarga tercinta yang senantiasa mendoakan dan memberikan dukungan мораil serta materiil.",
        "Rekan-rekan seperjuangan siswa SMK yang telah saling mendukung dan bertukar pengalaman selama kegiatan PKL berlangsung."
    ]
    for idx_p, item_text in enumerate(penghargaan_list, start=1):
        p_item = doc.add_paragraph()
        p_item.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_item.paragraph_format.left_indent = Inches(0.4)
        p_item.paragraph_format.space_after = Pt(2)
        p_item.add_run(f"{idx_p}. {item_text}")

    p_kp3 = doc.add_paragraph()
    p_kp3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kp3.paragraph_format.first_line_indent = Inches(0.4)
    p_kp3.paragraph_format.space_before = Pt(6)
    p_kp3.add_run(
        "Penulis menyadari bahwa dalam penyusunan laporan ini masih terdapat kekurangan karena keterbatasan pengetahuan dan pengalaman. "
        "Oleh karena itu, penulis sangat mengharapkan kritik dan saran yang membangun demi penyempurnaan laporan ini di masa mendatang. "
        "Semoga laporan ini dapat memberikan manfaat bagi pembaca, khususnya bagi adik-adik kelas."
    )

    # Titimangsa & Tanda Tangan Penulis di Kanan
    p_ttd_kp = doc.add_paragraph()
    p_ttd_kp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ttd_kp.paragraph_format.space_before = Pt(14)
    p_ttd_kp.paragraph_format.line_spacing = 1.15
    p_ttd_kp.add_run(f"{kota or 'Tempat'}, {datetime.now().strftime('%d %B %Y')}\nPenulis,\n\n\n\n\n")
    r_nm_penulis = p_ttd_kp.add_run(f"{nama_siswa.upper() if nama_siswa else 'NAMA PENULIS'}\n")
    r_nm_penulis.bold = True
    p_ttd_kp.add_run(f"NIS/NISN: {nis or '..............................'}")

    doc.add_page_break()

    # =========================================================================
    # 4. HALAMAN DAFTAR ISI (TABLE OF CONTENTS)
    # =========================================================================
    p_toc_hdr = doc.add_paragraph()
    p_toc_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_th = p_toc_hdr.add_run("DAFTAR ISI\n")
    run_th.bold = True
    run_th.font.size = Pt(14)

    # Bagian Awal (Angka Romawi)
    add_toc_line(doc, "HALAMAN JUDUL / SAMPUL", "i", is_bold=True)
    add_toc_line(doc, "LEMBAR PENGESAHAN", "ii", is_bold=True)
    add_toc_line(doc, "KATA PENGANTAR", "iii", is_bold=True)
    add_toc_line(doc, "DAFTAR ISI", "iv", is_bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Bab I
    add_toc_line(doc, "BAB I   PENDAHULUAN", "1", is_bold=True)
    add_toc_line(doc, "1.1 Latar Belakang Masalah", "1", indent_left=0.25)
    add_toc_line(doc, "1.2 Rumusan Masalah", "2", indent_left=0.25)
    add_toc_line(doc, "1.3 Batasan Masalah", "2", indent_left=0.25)
    add_toc_line(doc, "1.4 Tujuan PKL", "3", indent_left=0.25)
    add_toc_line(doc, "1.5 Manfaat PKL", "3", indent_left=0.25)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Bab II
    add_toc_line(doc, "BAB II  GAMBARAN UMUM PERUSAHAAN / DUDI", "4", is_bold=True)
    add_toc_line(doc, "2.1 Sejarah Singkat dan Profil Instansi", "4", indent_left=0.25)
    add_toc_line(doc, "2.2 Visi dan Misi Perusahaan", "5", indent_left=0.25)
    add_toc_line(doc, "2.3 Struktur Organisasi", "6", indent_left=0.25)
    add_toc_line(doc, "2.4 Tata Tertib dan Keselamatan Kerja (K3LH)", "7", indent_left=0.25)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Bab III
    add_toc_line(doc, "BAB III PELAKSANAAN DAN PEMBAHASAN TEKNIS", "8", is_bold=True)
    add_toc_line(doc, "3.1 Waktu dan Tempat Pelaksanaan", "8", indent_left=0.25)
    add_toc_line(doc, "3.2 Landasan Teori Singkat", "9", indent_left=0.25)
    add_toc_line(doc, "3.3 Alat dan Bahan Kerja", "11", indent_left=0.25)
    add_toc_line(doc, "3.4 Prosedur Pelaksanaan & Langkah-Langkah Teknis", "13", indent_left=0.25)
    add_toc_line(doc, "3.5 Hasil Pekerjaan dan Pengujian Sistem", "18", indent_left=0.25)
    add_toc_line(doc, "3.6 Kendala yang Dihadapi dan Solusi Pemecahan", "21", indent_left=0.25)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Bab IV
    add_toc_line(doc, "BAB IV  PENUTUP", "23", is_bold=True)
    add_toc_line(doc, "4.1 Kesimpulan", "23", indent_left=0.25)
    add_toc_line(doc, "4.2 Saran", "24", indent_left=0.25)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Bagian Akhir & Lampiran
    add_toc_line(doc, "DAFTAR PUSTAKA", "25", is_bold=True)
    add_toc_line(doc, f"LAMPIRAN 1: JURNAL / LOGBOOK MINGGUAN ({total_minggu} MINGGU)", "26", is_bold=True)
    add_toc_line(doc, "LAMPIRAN 2: FOTO DOKUMENTASI KEGIATAN & KERJA NYATA", "28", is_bold=True)

    doc.add_page_break()

    # =========================================================================
    # 5. KERANGKA LAPORAN (BAB I - BAB IV)
    # =========================================================================
    hdr_kerangka = doc.add_paragraph()
    hdr_kerangka.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_hk = hdr_kerangka.add_run("SISTEMATIKA & KERANGKA LAPORAN PKL\n")
    run_hk.bold = True
    run_hk.font.size = Pt(14)

    lines = ai_kerangka.split('\n')
    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue

        if re.match(r'^(#\s|BAB\s+[IVXLCDM]+)', line, re.IGNORECASE):
            clean_text = re.sub(r'^[#*_\s]+|[#*_\s]+$', '', line)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_text)
            run.bold = True
            run.font.size = Pt(13)

        elif re.match(r'^(##\s+|[A-Z]\.\s+|\d+\.\d+\s+)', line):
            clean_text = re.sub(r'^[#*_\s]+|[#*_\s]+$', '', line)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_text)
            run.bold = True
            run.font.size = Pt(12)

        elif line.startswith(('-', '*', '•')) or re.match(r'^\d+\.\s+', line):
            clean_text = re.sub(r'^[-*•\d.]+\s*', '', line)
            clean_text = re.sub(r'\*\*(.*?)\*\*', r'\1', clean_text)
            p = doc.add_paragraph(style='List Bullet')
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(clean_text)
            run.font.size = Pt(12)

        else:
            clean_text = re.sub(r'\*\*(.*?)\*\*', r'\1', line)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(clean_text)
            run.font.size = Pt(12)

    doc.add_page_break()

    # =========================================================================
    # 6. TABEL LOGBOOK KEGIATAN MINGGUAN (DINAMIS SESUAI TOTAL MINGGU)
    # =========================================================================
    hdr_logbook = doc.add_paragraph()
    hdr_logbook.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_log = hdr_logbook.add_run(f"JURNAL / LOGBOOK KEGIATAN MINGGUAN PKL ({durasi_text.upper()})\n")
    run_log.bold = True
    run_log.font.size = Pt(14)

    p_log_sub = doc.add_paragraph(f"Rekapitulasi tahapan pelaksanaan pekerjaan dan capaian kompetensi siswa selama {total_minggu} minggu:")
    p_log_sub.paragraph_format.space_after = Pt(8)

    headers = ["No", "Minggu", "Uraian Pekerjaan & Aktivitas Nyata", "Capaian / Hasil Belajar", "Paraf"]
    widths = [Inches(0.4), Inches(0.9), Inches(2.5), Inches(1.8), Inches(0.8)]

    table = doc.add_table(rows=total_minggu + 1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(10)
        shading = parse_xml(r'<w:shd {} w:fill="E0E0E0"/>'.format(nsdecls('w')))
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=80, right=80)

    # Isi Data Baris Sesuai Durasi
    for idx in range(total_minggu):
        row_idx = idx + 1
        row_cells = table.rows[row_idx].cells

        minggu_ke = idx + 1
        uraian = ""
        capaian = ""

        if idx < len(ai_logbook):
            entry = ai_logbook[idx]
            minggu_ke = entry.minggu_ke if hasattr(entry, 'minggu_ke') else entry.get('minggu_ke', idx + 1)
            uraian = entry.uraian_pekerjaan if hasattr(entry, 'uraian_pekerjaan') else entry.get('uraian_pekerjaan', '')
            capaian = entry.capaian_belajar if hasattr(entry, 'capaian_belajar') else entry.get('capaian_belajar', '')

        # Kolom No
        row_cells[0].text = str(idx + 1)
        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row_cells[0].paragraphs[0].runs[0].font.size = Pt(10)

        # Kolom Minggu
        row_cells[1].text = f"Minggu {minggu_ke}"
        row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row_cells[1].paragraphs[0].runs[0].font.size = Pt(10)

        # Kolom Uraian Pekerjaan
        row_cells[2].text = uraian
        row_cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        row_cells[2].paragraphs[0].runs[0].font.size = Pt(9.5)

        # Kolom Capaian
        row_cells[3].text = capaian
        row_cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        row_cells[3].paragraphs[0].runs[0].font.size = Pt(9.5)

        # Kolom Paraf
        row_cells[4].text = ""

        for col_i in range(5):
            row_cells[col_i].width = widths[col_i]
            set_cell_margins(row_cells[col_i], top=120, bottom=120, left=70, right=70)

    tblPr = table._tbl.tblPr
    borders = parse_xml(
        r'<w:tblBorders {} >'
        r'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'</w:tblBorders>'.format(nsdecls('w'))
    )
    tblPr.append(borders)

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream


# --- UI UTAMA STREAMLIT ---
st.title("🎓 Platform Pembuatan Laporan & Logbook PKL SMK")
st.caption("Menghasilkan Dokumen Word Lengkap: Cover Resmi, Lembar Pengesahan, Kata Pengantar, Daftar Isi (TOC), Kerangka Laporan AI, dan Logbook Mingguan.")

# Cek apakah API key tersedia di Streamlit Secrets (server) atau input manual
if "GEMINI_API_KEY" in st.secrets:
    api_key = st.secrets["GEMINI_API_KEY"]
else:
    api_key = st.sidebar.text_input("Gemini API Key", type="password")
st.sidebar.markdown("---")
st.sidebar.subheader("🖼️ Upload Logo Sekolah")
uploaded_logo = st.sidebar.file_uploader(
    "Unggah Logo Sekolah (PNG / JPG)",
    type=["png", "jpg", "jpeg"],
    help="Otomatis ditempatkan di halaman sampul dokumen Word."
)
if uploaded_logo:
    st.sidebar.image(uploaded_logo, caption="Pratinjau Logo", width=150)

with st.form("form_pkl"):
    st.subheader("1. Data Identitas, Durasi & Sampul")
    c1, c2, c3 = st.columns(3)
    with c1:
        nama_siswa = st.text_input("Nama Siswa", placeholder="contoh: Muhammad Rizky")
        nis = st.text_input("NIS / NISN", placeholder="contoh: 212210045")
    with c2:
        sekolah = st.text_input("Nama Sekolah", placeholder="contoh: SMK Negeri 1 Masbagik")
        jurusan = st.selectbox(
            "Kompetensi Keahlian",
            [
                "Teknik Komputer dan Jaringan (TKJ)",
                "Rekayasa Perangkat Lunak (RPL)",
                "Teknik Kendaraan Ringan (TKR)",
                "Teknik Bisnis Sepeda Motor (TBSM)",
                "Akuntansi & Keuangan Lembaga (AKL)",
                "Otomatisasi & Tata Kelola Perkantoran (OTKP)",
                "Bisnis Daring & Pemasaran (BDP)",
                "Lainnya"
            ]
        )
    with c3:
        durasi_pilihan = st.selectbox(
            "Durasi Pelaksanaan PKL",
            [
                "3 Bulan (12 Minggu) - Reguler",
                "6 Bulan (24 Minggu) - Kurikulum Merdeka / SMK PK"
            ]
        )
        total_minggu = 12 if "3 Bulan" in durasi_pilihan else 24
        durasi_label = "3 Bulan" if total_minggu == 12 else "6 Bulan"

        kota = st.text_input("Kota / Kabupaten Sekolah", placeholder="contoh: Lombok Timur")
        curr_year = datetime.now().year
        tahun_ajaran = st.text_input("Tahun Ajaran", value=f"{curr_year}/{curr_year + 1}")

    st.subheader("2. Instansi & Pejabat Pengesahan")
    p1, p2, p3 = st.columns(3)
    with p1:
        tempat_dudi = st.text_input("Nama DUDI / Tempat PKL", placeholder="contoh: PT Cyber Network")
    with p2:
        nama_pembimbing_sekolah = st.text_input("Guru Pembimbing Sekolah", placeholder="Nama & Gelar")
        nama_pembimbing_dudi = st.text_input("Pembimbing DU/DI Lapangan", placeholder="Nama Pembimbing Industri")
    with p3:
        nama_kepsek = st.text_input("Kepala Sekolah", placeholder="Nama Kepala Sekolah & Gelar")

    st.subheader("3. Rincian Teknis Laporan & Praktik")
    judul_pkl = st.text_area(
        "Judul Laporan PKL",
        placeholder="contoh: Rancang Bangun dan Konfigurasi Jaringan Hotspot Berbasis RADIUS Server pada PT Cyber Network",
        help=f"Judul spesifik memudahkan AI membagi alur kerja selama {total_minggu} minggu secara realistis."
    )
    catatan_tambahan = st.text_input(
        "Fokus Alat, Perangkat, atau SOP Khusus (Opsional)",
        placeholder="contoh: Mikrotik RB750Gr3, Winbox, FreeRADIUS, daloRADIUS"
    )

    submitted = st.form_submit_button(f"🚀 Buat Dokumen Lengkap ({total_minggu} Minggu Logbook)", use_container_width=True)

if submitted:
    if not api_key:
        st.error("Silakan masukkan Gemini API Key di bilah samping.")
    elif not judul_pkl.strip():
        st.warning("Judul laporan PKL tidak boleh kosong.")
    else:
        with st.spinner(f"AI sedang merancang sistematika bab dan memetakan jadwal logbook {total_minggu} minggu..."):
            try:
                client = genai.Client(api_key=api_key)

                if total_minggu == 12:
                    pembagian_tahapan = """
                    - Minggu 1-2: Orientasi DUDI, K3, pengenalan lingkungan dan observasi sistem awal.
                    - Minggu 3-4: Penyiapan alat/bahan, analisis kebutuhan, perancangan skema/topologi.
                    - Minggu 5-8: Pelaksanaan inti pekerjaan teknis / konfigurasi / praktik utama sesuai judul.
                    - Minggu 9-10: Pengujian sistem, uji performa, troubleshooting kendala, pemeliharaan.
                    - Minggu 11-12: Evaluasi hasil kerja, dokumentasi, penyusunan laporan, penarikan siswa.
                    """
                else:
                    pembagian_tahapan = """
                    - Minggu 1-3: Orientasi DUDI, pengenalan budaya kerja industri, K3LH, observasi menyeluruh sistem & SOP operasional.
                    - Minggu 4-6: Penyiapan alat, analisis kelayakan, audit infrastruktur/proses eksisting, perancangan solusi teknis.
                    - Minggu 7-14: Pelaksanaan inti pekerjaan teknis secara mendalam, implementasi proyek/fitur/instalasi bertahap, konfigurasi lanjut.
                    - Minggu 15-18: Operasional harian, pemeliharaan rutin (preventive maintenance), penanganan kasus gangguan lapangan (troubleshooting).
                    - Minggu 19-21: Uji reliabilitas sistem, optimasi performa, quality control / audit hasil kerja bersama mentor industri.
                    - Minggu 22-24: Penyusunan dokumentasi teknis akhir, presentasi hasil di depan manajemen DUDI, evaluasi, penarikan siswa.
                    """

                prompt = f"""
                Kamu adalah guru pembimbing dan koordinator PKL SMK yang profesional.
                Berdasarkan data berikut:
                - Judul Laporan: {judul_pkl}
                - Jurusan: {jurusan}
                - Tempat DUDI: {tempat_dudi or 'Instansi DUDI Terkait'}
                - Durasi PKL: {durasi_label} ({total_minggu} Minggu)
                - Fokus Alat/Kegiatan: {catatan_tambahan or 'Disesuaikan dengan judul'}

                Tugasmu menghasilkan objek data dengan 2 properti:
                1. `kerangka_laporan`: Kerangka lengkap laporan berstruktur baku:
                   BAB I PENDAHULUAN (Latar Belakang, Rumusan Masalah, Batasan Masalah, Tujuan, Manfaat)
                   BAB II GAMBARAN UMUM PERUSAHAAN (Sejarah, Visi Misi, Struktur, K3)
                   BAB III PELAKSANAAN DAN PEMBAHASAN TEKNIS (Waktu & Tempat, Landasan Teori, Alat & Bahan, Prosedur Teknis Langkah demi Langkah, Hasil & Pengujian, Kendala & Solusi)
                   BAB IV PENUTUP (Kesimpulan, Saran Sekolah, Saran DUDI)
                   DAFTAR PUSTAKA & LAMPIRAN

                2. `logbook_mingguan`: Buat TEPAT {total_minggu} entri kegiatan mingguan (minggu 1 sampai minggu {total_minggu}) yang realistis dan mendalam dengan pedoman tahapan:
                {pembagian_tahapan}

                Pastikan istilah teknis, alat, dan prosedur sangat relevan dengan bidang keahlian {jurusan}.
                """

                response = client.models.generate_content(
                    model="gemini-3.8-flash",
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        temperature=0.3,
                        response_mime_type="application/json",
                        response_schema=LaporanPKLOutput
                    )
                )

                result_data = json.loads(response.text)
                ai_kerangka = result_data.get("kerangka_laporan", "")
                ai_logbook = result_data.get("logbook_mingguan", [])

                # Generate File Word Lengkap (.docx)
                docx_file = create_word_document(
                    nama_siswa=nama_siswa,
                    nis=nis,
                    sekolah=sekolah,
                    kota=kota,
                    tahun_ajaran=tahun_ajaran,
                    judul_text=judul_pkl,
                    ai_kerangka=ai_kerangka,
                    ai_logbook=ai_logbook,
                    jurusan=jurusan,
                    tempat_dudi=tempat_dudi,
                    durasi_text=durasi_label,
                    total_minggu=total_minggu,
                    nama_pembimbing_sekolah=nama_pembimbing_sekolah,
                    nama_pembimbing_dudi=nama_pembimbing_dudi,
                    nama_kepsek=nama_kepsek,
                    logo_file=uploaded_logo
                )

                st.success(f"✅ Dokumen laporan lengkap (Cover, Pengesahan, Kata Pengantar, Daftar Isi, Bab I-IV, Logbook {total_minggu} Minggu) berhasil dibuat!")

                c_dl1, c_dl2 = st.columns(2)
                with c_dl1:
                    st.download_button(
                        label="📄 Unduh Dokumen Lengkap (.docx)",
                        data=docx_file,
                        file_name=f"Laporan_PKL_{durasi_label}_{nama_siswa.replace(' ', '_') if nama_siswa else 'Siswa'}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True
                    )
                with c_dl2:
                    st.download_button(
                        label="📝 Unduh Kerangka Teks (.txt)",
                        data=ai_kerangka,
                        file_name=f"Kerangka_PKL_{durasi_label}.txt",
                        mime="text/plain",
                        use_container_width=True
                    )

                # Tabs Pratinjau
                tab_logbook, tab_kerangka = st.tabs([f"📋 Pratinjau Logbook ({total_minggu} Minggu)", "📖 Pratinjau Kerangka Bab I - IV"])
                
                with tab_logbook:
                    log_rows = []
                    for item in ai_logbook:
                        log_rows.append({
                            "Minggu": f"Minggu ke-{item['minggu_ke']}",
                            "Uraian Pekerjaan / Aktivitas Nyata": item["uraian_pekerjaan"],
                            "Capaian / Hasil Belajar": item["capaian_belajar"]
                        })
                    st.table(log_rows)

                with tab_kerangka:
                    st.markdown(ai_kerangka)

            except Exception as e:
                st.error(f"Gagal memproses dokumen: {str(e)}")
